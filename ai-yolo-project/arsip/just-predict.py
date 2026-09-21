# predict_yolo_autopilot.py

import os
import json
from pathlib import Path
from datetime import datetime
import shutil
from tqdm import tqdm

import pandas as pd
from docx import Document
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet

import matplotlib.pyplot as plt

from ultralytics import YOLO
from PIL import Image

# ================= CONFIG (terpusat di config.py) =================
from config import (BEST_MODEL, SCAN_DIR, HIGH_CONF_DIR, MEDIUM_CONF_DIR,
                    LOW_CONF_DIR, ERROR_IMAGES_DIR, REPORT_DIR,
                    HIGH_CONF_THRESHOLD, LOW_CONF_THRESHOLD, PROJECT_ROOT)

MODEL_PATH = str(BEST_MODEL)

BASE = SCAN_DIR.parent
PROJECT = PROJECT_ROOT

SCAN = SCAN_DIR

FOLDERS = {
    "high": HIGH_CONF_DIR,
    "medium": MEDIUM_CONF_DIR,
    "low": LOW_CONF_DIR,
    "error": ERROR_IMAGES_DIR,
}

HIGH_TH = HIGH_CONF_THRESHOLD
LOW_TH = LOW_CONF_THRESHOLD
HISTORY_FILE = REPORT_DIR / "history.json"

for p in list(FOLDERS.values()) + [REPORT_DIR]:
    p.mkdir(parents=True, exist_ok=True)

# ================= MODEL =================
model = YOLO(MODEL_PATH)

# ================= UTILS =================
def is_valid(img):
    try:
        Image.open(img).verify()
        return True
    except:
        return False


def safe_move(src, dst_folder, class_name=""):
    dst_folder = Path(dst_folder) / class_name if class_name else Path(dst_folder)
    dst_folder.mkdir(parents=True, exist_ok=True)

    dst = dst_folder / Path(src).name

    i = 1
    while dst.exists():
        dst = dst_folder / f"{dst.stem}_dup{i}{dst.suffix}"
        i += 1

    shutil.move(str(src), str(dst))


def load_history():
    if HISTORY_FILE.exists():
        return json.loads(HISTORY_FILE.read_text())
    return []


def save_history(data):
    history = load_history()
    history.append(data)
    HISTORY_FILE.write_text(json.dumps(history, indent=2))


# ================= CREATE RUN REPORT FOLDER =================
now = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")

RUN_REPORT_DIR = REPORT_DIR / now
RUN_REPORT_DIR.mkdir(parents=True, exist_ok=True)

txt_path = RUN_REPORT_DIR / "report.txt"
excel_path = RUN_REPORT_DIR / "report.xlsx"
word_path = RUN_REPORT_DIR / "report.docx"
pdf_path = RUN_REPORT_DIR / "report.pdf"
graph_path = RUN_REPORT_DIR / "trend.png"

# ================= LOAD DATA =================
images = list(SCAN.rglob("*.*"))
print(f"[INFO] Images found: {len(images)}")

stats = {
    "high": 0,
    "medium": 0,
    "low": 0,
    "error": 0
}

category_count = {}

# ================= PROCESS =================
for img in tqdm(images, desc="AI Processing"):

    img = Path(img)

    if not is_valid(img):
        safe_move(img, FOLDERS["error"])
        stats["error"] += 1
        continue

    try:
        r = model.predict(str(img), imgsz=224, conf=0.5)[0]

        idx = r.probs.top1
        conf = float(r.probs.top1conf)
        name = r.names[idx]

        category_count[name] = category_count.get(name, 0) + 1

        if conf >= HIGH_TH:
            safe_move(img, FOLDERS["high"], name)
            stats["high"] += 1

        elif conf < LOW_TH:
            safe_move(img, FOLDERS["low"], name)
            stats["low"] += 1

        else:
            safe_move(img, FOLDERS["medium"], name)
            stats["medium"] += 1

    except:
        safe_move(img, FOLDERS["error"])
        stats["error"] += 1


# ================= HISTORY =================
history = load_history()

prev = history[-1] if history else None

current = {
    "time": now,
    "total": len(images),
    "stats": stats
}

save_history(current)


def diff(key):
    if not prev:
        return 0
    return stats[key] - prev["stats"].get(key, 0)


insight = {
    "high_change": diff("high"),
    "medium_change": diff("medium"),
    "low_change": diff("low"),
}

# ================= TXT REPORT =================
txt_path.write_text(f"""
YOLO REPORT {now}

TOTAL: {len(images)}

HIGH: {stats['high']}
MEDIUM: {stats['medium']}
LOW: {stats['low']}
ERROR: {stats['error']}

INSIGHT:
HIGH change: {insight['high_change']}
MEDIUM change: {insight['medium_change']}
LOW change: {insight['low_change']}
""")


# ================= EXCEL =================
df = pd.DataFrame([{
    "time": now,
    "total": len(images),
    **stats,
    **insight
}])

df.to_excel(excel_path, index=False)


# ================= WORD =================
doc = Document()
doc.add_heading("YOLO REPORT", 0)

doc.add_paragraph(f"Time: {now}")
doc.add_paragraph(f"Total: {len(images)}")

doc.add_heading("Stats", level=1)
for k, v in stats.items():
    doc.add_paragraph(f"{k}: {v}")

doc.add_heading("AI Improvement", level=1)
for k, v in insight.items():
    doc.add_paragraph(f"{k}: {v}")

doc.save(word_path)


# ================= PDF =================
pdf_doc = SimpleDocTemplate(str(pdf_path))
style = getSampleStyleSheet()

content = [
    Paragraph("YOLO REPORT", style["Title"]),
    Spacer(1, 12),
    Paragraph(f"Time: {now}", style["Normal"]),
    Paragraph(f"Total: {len(images)}", style["Normal"]),
    Spacer(1, 12),
]

for k, v in stats.items():
    content.append(Paragraph(f"{k}: {v}", style["Normal"]))

pdf_doc.build(content)


# ================= GRAPH =================
if len(history) > 1:
    times = [h["time"] for h in history]
    highs = [h["stats"]["high"] for h in history]

    plt.figure()
    plt.plot(times, highs)
    plt.xticks(rotation=45)
    plt.title("AI Improvement Trend (HIGH confidence)")
    plt.tight_layout()
    plt.savefig(graph_path)
    plt.close()


# ================= FINAL =================
print("\n🔥 DONE AI AUTOPILOT SYSTEM")
print(f"REPORT FOLDER : {RUN_REPORT_DIR}")
print(f"TXT   : {txt_path}")
print(f"EXCEL : {excel_path}")
print(f"WORD  : {word_path}")
print(f"PDF   : {pdf_path}")
print(f"GRAPH : {graph_path}")